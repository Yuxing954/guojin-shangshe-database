#!/usr/bin/env python3
"""Import Hotel Home holiday operating comparisons from source workbooks."""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import json
import re
from pathlib import Path

from openpyxl import load_workbook


SOURCES = {
    "spring-festival": ("春节", "2025&2026春节假期.xlsx"),
    "qingming": ("清明", "2025&2026清明假期.xlsx"),
    "mid-autumn": ("中秋", "2026&2024中秋假期.xlsx"),
    "dragon-boat": ("端午", "2026&2025端午.xlsx"),
    "national-day": ("国庆", "2026&2025国庆假期.xlsx"),
}
DIMENSIONS = {"分档位": "hotel_tier", "分线级": "city_tier"}
FIELDS = [
    "holiday_id", "holiday_name", "dimension", "segment", "period_day",
    "current_year", "baseline_year", "current_date", "baseline_date", "day_label",
    "occupancy_rate", "adr", "revpar", "baseline_occupancy_rate", "baseline_adr", "baseline_revpar",
    "source_file", "source_sheet", "source_row",
]
DATE_RE = re.compile(r"^(\d{4})[/-](\d{1,2})[/-](\d{1,2})")


def date_parts(value):
    if isinstance(value, (dt.datetime, dt.date)):
        return value.date() if isinstance(value, dt.datetime) else value, ""
    if isinstance(value, str):
        match = DATE_RE.match(value.strip())
        if match:
            date = dt.date(*(int(part) for part in match.groups()))
            label = re.search(r"[（(]([^）)]+)[）)]", value)
            return date, label.group(1).strip() if label else ""
    return None, ""


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def import_workbook(path: Path, holiday_id: str, holiday_name: str):
    workbook = load_workbook(path, data_only=True, read_only=True)
    output, coverage = [], {}
    for sheet in workbook.worksheets:
        dimension = DIMENSIONS.get(sheet.title)
        if not dimension:
            continue
        current_date = baseline_date = None
        day_label = ""
        day_index = 0
        sheet_rows = 0
        for row_number, row in enumerate(sheet.iter_rows(min_row=3, values_only=True), start=3):
            if len(row) < 10 or not row[1] or not row[6]:
                continue
            parsed_current, parsed_label = date_parts(row[0])
            parsed_baseline, _ = date_parts(row[5])
            if parsed_current:
                current_date = parsed_current
                day_label = parsed_label
                day_index += 1
            if parsed_baseline:
                baseline_date = parsed_baseline
            if current_date is None or baseline_date is None:
                raise ValueError(f"{path.name} / {sheet.title} row {row_number}: missing date mapping")
            if row[1] != row[6]:
                raise ValueError(f"{path.name} / {sheet.title} row {row_number}: segment labels differ")
            record = {
                "holiday_id": holiday_id,
                "holiday_name": holiday_name,
                "dimension": dimension,
                "segment": str(row[1]).strip(),
                "period_day": day_index,
                "current_year": current_date.year,
                "baseline_year": baseline_date.year,
                "current_date": current_date.isoformat(),
                "baseline_date": baseline_date.isoformat(),
                "day_label": day_label,
                "occupancy_rate": row[2],
                "adr": row[3],
                "revpar": row[4],
                "baseline_occupancy_rate": row[7],
                "baseline_adr": row[8],
                "baseline_revpar": row[9],
                "source_file": path.name,
                "source_sheet": sheet.title,
                "source_row": row_number,
            }
            for field in ("occupancy_rate", "adr", "revpar", "baseline_occupancy_rate", "baseline_adr", "baseline_revpar"):
                if record[field] is not None and not isinstance(record[field], (int, float)):
                    raise ValueError(f"{path.name} / {sheet.title} row {row_number}: invalid {field}")
            output.append(record)
            sheet_rows += 1
        coverage[sheet.title] = {
            "dimension": dimension,
            "rows": sheet_rows,
            "periodDays": day_index,
            "currentStart": min((r["current_date"] for r in output if r["source_sheet"] == sheet.title), default=None),
            "currentEnd": max((r["current_date"] for r in output if r["source_sheet"] == sheet.title), default=None),
            "baselineStart": min((r["baseline_date"] for r in output if r["source_sheet"] == sheet.title), default=None),
            "baselineEnd": max((r["baseline_date"] for r in output if r["source_sheet"] == sheet.title), default=None),
        }
    workbook.close()
    if not output:
        raise ValueError(f"No hotel records found in {path.name}")
    return output, coverage


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source_dir", type=Path, help="Folder containing the five Hotel Home workbooks")
    parser.add_argument("data_dir", type=Path, help="Repository data directory")
    args = parser.parse_args()

    records, sources, coverage = [], [], {}
    for holiday_id, (holiday_name, filename) in SOURCES.items():
        path = args.source_dir / filename
        if not path.is_file():
            raise FileNotFoundError(path)
        rows, sheets = import_workbook(path, holiday_id, holiday_name)
        records.extend(rows)
        sources.append({"file": filename, "sha256": sha256(path), "rows": len(rows)})
        coverage[holiday_id] = {"name": holiday_name, "sheets": sheets}

    keys = [(r["holiday_id"], r["dimension"], r["segment"], r["period_day"]) for r in records]
    if len(keys) != len(set(keys)):
        raise ValueError("Duplicate holiday/dimension/segment/day keys found")
    records.sort(key=lambda r: (r["current_date"], r["holiday_id"], r["dimension"], r["period_day"], r["segment"]))

    args.data_dir.mkdir(parents=True, exist_ok=True)
    csv_path = args.data_dir / "hotel_holiday_daily.csv"
    with csv_path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=FIELDS, extrasaction="raise")
        writer.writeheader()
        writer.writerows(records)

    imported_at = dt.datetime.now(dt.timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")
    manifest = {
        "updatedAt": dt.date.today().isoformat(),
        "importedAt": imported_at,
        "source": "酒店之家",
        "file": "data/hotel_holiday_daily.csv",
        "rowCount": len(records),
        "key": ["holiday_id", "dimension", "segment", "period_day"],
        "comparisonAlignment": "holiday day index; current and baseline dates are retained as supplied",
        "sources": sources,
        "coverage": coverage,
    }
    (args.data_dir / "hotel-holiday-import-record.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(json.dumps({"file": str(csv_path), "rows": len(records), "coverage": coverage}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()

