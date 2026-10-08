import argparse, csv, json, re, warnings
from pathlib import Path
import openpyxl
warnings.filterwarnings('ignore', category=UserWarning, module='openpyxl')
ROOT=Path(__file__).resolve().parents[1]
parser=argparse.ArgumentParser(description='Import duty-free holiday observations from the research workbook.')
parser.add_argument('source',type=Path)
parser.add_argument('--source-version',required=True)
parser.add_argument('--imported-at',required=True)
args=parser.parse_args()
SOURCE=args.source
w=openpyxl.load_workbook(SOURCE,data_only=True)
f=openpyxl.load_workbook(SOURCE,data_only=False)
records=[]
def value(sheet,col,row):
    v=w[sheet].cell(row,col).value
    if isinstance(v,(float,int)) and not isinstance(v,bool):return v
    if isinstance(v,str) and re.fullmatch(r'[+-]?(?:\d+(?:\.\d*)?|\.\d+)',v.strip()):return float(v.strip())
    return None
def record(sheet,col,header,period,days,total,kind='holiday',label=None,year=None):
    s=w[sheet]; h=str(s.cell(header,col).value or '')
    match=re.search(r'20\d{2}',h)
    year=year or (int(match.group()) if match else None)
    if not year:return
    holiday=label or h.replace(str(year),'').replace('假期','').strip()
    refs={}
    def read(key,row):
        cell=s.cell(row,col);refs[key]={'cell':cell.coordinate,'value':cell.value,'formula':f[sheet][cell.coordinate].value if f[sheet][cell.coordinate].data_type=='f' else None}
        return value(sheet,col,row)
    sales=read('sales',total);shoppers=read('shoppers',total+1);spend=read('spend',total+2)
    count=read('days',days) if days else None
    period_text=str(s.cell(period,col).value or '')
    refs['period']={'cell':s.cell(period,col).coordinate,'value':period_text}
    notes=[]
    if sheet=='国庆中秋假期' and year==2022:
        notes.append('原表日期为9.29-10.6，与年份对应关系待核验；未推定统计天数。')
    pending=sales is None and shoppers is None
    daily_sales=sales/count if sales is not None and count else None
    daily_shoppers=shoppers/count if shoppers is not None and count else None
    if kind=='holiday' and header in (8,16):
        base=header+3
        src_daily=read('source_daily_sales',base); src_people=read('source_daily_shoppers',base+1);src_spend=read('source_summary_spend',base+2)
        for key,a,b in [('日均销售额',src_daily,daily_sales),('日均购物人次',src_people,daily_shoppers),('客单价',src_spend,spend)]:
            if a is not None and b is not None and abs(a-b)>max(.01,abs(b)*.001):notes.append(f'原表汇总{key}与总额区域不一致，采用总额区域对应口径，原值保留在来源明细。')
    if pending:
        notes.append('原表总额空缺；日均及同比公式的0或-100%不作为已披露数据。')
        spend=None
    formula_inputs=any(refs[k]['formula'] for k in ('sales','shoppers'))
    quality='原表公式值' if formula_inputs else '原表录入值'
    r={'id':f'{year}-{holiday}-{kind}-{sheet}-{col}-{total}', 'year':year,'holiday':holiday,'kind':kind,'period':period_text,'days':count,'sales':sales,'shoppers':shoppers,'spend':spend,'daily_sales':daily_sales,'daily_shoppers':daily_shoppers,'status':'pending' if pending else 'available','quality':quality,'sheet':sheet,'source_cells':refs,'notes':notes,'daily_yoy':None,'shoppers_daily_yoy':None,'spend_yoy':None,'yoy_basis':None}
    if kind=='holiday' and not pending:
        for c in range(col+1,s.max_column+1):
            comp=str(s.cell(period,c).value or '')
            if comp==f'{str(year)[2:]}vs{str(year-1)[2:]}':
                for key,row in [('daily_yoy',header+3),('shoppers_daily_yoy',header+4),('spend_yoy',header+5)]:
                    v=value(sheet,c,row)
                    if v is not None:
                        r[key]=v*100
                        refs[key]={'cell':s.cell(row,c).coordinate,'value':v,'formula':f[sheet].cell(row,c).value if f[sheet].cell(row,c).data_type=='f' else None}
                r['yoy_basis']='原表同比栏（当年公布口径）'
                break
    records.append(r)
    return r
# Whole-holiday totals are kept separate from calendar-day/partial-period figures.
for sheet,cols,total in [('元旦假期',range(3,10),42),('春节假期',range(3,8),17),('五一假期',range(3,10),28),('清明假期',[3,5],22),('端午假期',range(3,6),22)]:
    for col in cols:
        r=record(sheet,col,8,9,10,total)
        if sheet=='元旦假期' and r['year']==2025:
            # The holiday lasts one day. Rows 42-44 are the Jan 1-3 calendar window.
            for key,row in [('sales',16),('shoppers',17),('spend',18)]:
                r[key]=value(sheet,col,row);r['source_cells'][key]={'cell':w[sheet].cell(row,col).coordinate,'value':r[key],'formula':f[sheet].cell(row,col).value if f[sheet].cell(row,col).data_type=='f' else None}
            r['daily_sales']=r['sales'];r['daily_shoppers']=r['shoppers'];r['notes']=['2025年元旦法定假期仅1天；1.1-1.3同日历窗口另列在阶段数据中。']
for col in [3,5]:record('国庆中秋假期',col,8,9,10,30,label='中秋')
for col in range(3,7):
    r=record('国庆中秋假期',col,16,17,18,34)
    r['comparison_group']='国庆及中秋国庆'
record('国庆中秋假期',7,28,33,None,34,label='国庆',year=2022)['comparison_group']='国庆及中秋国庆'
# Preserve summary-source prior-year rates, without guessing missing prior-year amounts.
year=None
for row in range(2,12):
    year=value('主要假期',1,row) or year
    holiday=w['主要假期'].cell(row,2).value
    rr=next((r for r in records if r['kind']=='holiday' and r['year']==year and r['holiday'].replace('国庆中秋','中秋国庆')==holiday),None)
    v=value('主要假期',4,row)
    if rr and v is not None and rr['daily_yoy'] is None:
        rr['daily_yoy']=v*100;rr['yoy_basis']='主要假期表同比（当年公布口径）';rr['source_cells']['daily_yoy']={'sheet':'主要假期','cell':f'D{row}','value':v,'formula':f['主要假期'].cell(row,4).value}
# Additional actual observations, not added to whole-holiday totals.
for col in [3,4]:
    for row,days,label in [(16,1,'元旦首日'),(21,1,'元旦第二日'),(26,1,'元旦第三日'),(36,7,'元旦首周'),(42,3,'元旦同日历三天'),(47,2,'元旦前两天')]:
        r=record('元旦假期',col,8,row-1,None,row,'detail',label)
        if r and r['status']=='available':r['days']=days;r['daily_sales']=r['sales']/days if r['sales'] is not None else None;r['daily_shoppers']=r['shoppers']/days if r['shoppers'] is not None else None;r['notes'].append('阶段数据不可与完整假期加总。');r['holiday_group']='元旦'
        else:records.remove(r)
for col in [3,4]:
    for row,days,label in [(24,1,'春节首日'),(29,3,'春节前三天'),(34,4,'春节前四天'),(39,5,'春节前五天'),(45,None,'春运期间')]:
        r=record('春节假期',col,8,row-2 if row==24 else row-1,None,row,'detail',label)
        r['days']=days;r['daily_sales']=r['sales']/days if r['sales'] is not None and days else None;r['daily_shoppers']=r['shoppers']/days if r['shoppers'] is not None and days else None;r['holiday_group']='春节';r['notes'].append('阶段数据不可与完整假期加总。')
for col in [5,8,9]:
    r=record('五一假期',col,8,16,None,17,'detail','五一首日');r['days']=1;r['daily_sales']=r['sales'];r['daily_shoppers']=r['shoppers'];r['holiday_group']='五一'
# Recompute the type after the New Year one-day source replacement.
for r in records:r['quality']='原表公式值' if any(r['source_cells'][k].get('formula') for k in ['sales','shoppers']) else '原表录入值'
# Compute only missing same-holiday, consecutive-year comparisons; source rates take precedence.
for r in records:
    if r['kind']!='holiday' or r['status']!='available':continue
    group=r.get('comparison_group',r['holiday'])
    prev=next((p for p in records if p['kind']=='holiday' and p['status']=='available' and p['year']==r['year']-1 and p.get('comparison_group',p['holiday'])==group),None)
    if prev:
        for field,key in [('daily_sales','daily_yoy'),('daily_shoppers','shoppers_daily_yoy'),('spend','spend_yoy')]:
            if r[key] is None and r[field] is not None and prev[field]:r[key]=(r[field]/prev[field]-1)*100;r['yoy_basis']='缺失同比按本表相同假期上一年计算；已有同比沿用原表'
data={'schemaVersion':1,'sourceFile':SOURCE.name,'sourceVersion':args.source_version,'importedAt':args.imported_at,'source':'海口海关（Excel内标注），国金证券整理；本次按用户提供原表录入，未独立回查公告','notes':['金额单位：亿元；购物人次：万人次；客单价：元。','日均=同一期间总额/原表统计天数；不同假期天数比较优先看日均。','原表公式反推的历史数值保留标识，不当作独立披露；未发布或未填写值保持缺失。','阶段数据与假期总额存在重叠，不能加总。','原表假设测算区域未作为实际数据导入。'],'records':records}
# Retain independently sourced public updates when refreshing the older workbook.
existing_path=ROOT/'data/dutyfree/holidays.json'
if existing_path.exists():
    existing=json.loads(existing_path.read_text(encoding='utf-8'))
    public={r['id']:r for r in existing.get('records',[]) if r.get('web_source')}
    if public:
        data['records']=[public.pop(r['id'],r) for r in records]+list(public.values())
        records=data['records']
        data['updatedAt']=existing.get('updatedAt',args.imported_at)
        data['source']=existing['source']
        for r in records:
            if r.get('comparison_group')=='国庆及中秋国庆':r['comparison_group']='国庆'
            if r['holiday']=='国庆中秋':r['holiday']='中秋国庆'
(ROOT/'data/dutyfree/holidays.json').write_text(json.dumps(data,ensure_ascii=False,indent=2,allow_nan=False)+'\n',encoding='utf-8')
fields=['year','holiday','daily_sales_cny_100m','yoy_pct','source','period','days','sales_cny_100m','shoppers_10k','spend_per_shopper_cny','daily_shoppers_10k','status','quality','source_sheet','source_sales_cell']
with (ROOT/'data/dutyfree_holiday.csv').open('w',encoding='utf-8',newline='') as out:
    writer=csv.DictWriter(out,fieldnames=fields,lineterminator='\n');writer.writeheader()
    for r in records:
        if r['kind']!='holiday':continue
        writer.writerow(dict(zip(fields,[r['year'],r['holiday'],r['daily_sales'],r['daily_yoy'],r.get('web_source',{}).get('url',SOURCE.name),r['period'],r['days'],r['sales'],r['shoppers'],r['spend'],r['daily_shoppers'],r['status'],r['quality'],r['sheet'],r['source_cells'].get('sales',{}).get('cell','')])))
print(json.dumps({'total':len(records),'holidays':sum(r['kind']=='holiday' for r in records),'available_holidays':sum(r['kind']=='holiday' and r['status']=='available' for r in records),'pending':[f"{r['year']}{r['holiday']}" for r in records if r['status']=='pending']},ensure_ascii=False))
