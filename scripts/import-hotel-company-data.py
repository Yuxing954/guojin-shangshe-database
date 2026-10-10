"""Prepare source-traceable hotel company observations; never overwrites source files."""
from pathlib import Path
import argparse,csv,datetime,hashlib,json,re,math,collections
import openpyxl

METRICS={
 'hotels':('期末酒店数','家','规模','stock'),'rooms':('期末客房数','间','规模','stock'),
 'pipeline':('待开业酒店数','家','规模','stock'),'opened':('新开酒店数','家','拓店','flow'),
 'closed':('退出酒店数','家','拓店','flow'),'to_pipeline':('开业转筹建','家','拓店','flow'),'net_opened':('净增酒店数','家','拓店','flow'),
 'adr':('ADR','元/间','经营','average'),'occ':('入住率','%','经营','average'),'revpar':('RevPAR','元/间','经营','average'),
 'revenue':('营业收入','百万元','财务','flow'),'hotel_revenue':('酒店业务收入','百万元','财务','flow'),
 'managed_revenue':('加盟及管理收入','百万元','财务','flow'),'leased_revenue':('直营收入','百万元','财务','flow'),
 'retail_revenue':('零售收入','百万元','财务','flow'),'retail_other_revenue':('零售及其他收入','百万元','财务','flow'),
 'other_revenue':('其他收入','百万元','财务','flow'),'operating_profit':('经营利润','百万元','财务','flow'),
 'net_profit':('净利润','百万元','财务','flow'),'parent_profit':('归母净利润','百万元','财务','flow'),
 'adjusted_profit':('经调整净利润','百万元','财务','flow'),'adjusted_ebitda':('经调整EBITDA','百万元','财务','flow'),
 'operating_cf':('经营现金流','百万元','现金流','flow'),'investing_cf':('投资现金流','百万元','现金流','flow'),
 'financing_cf':('融资现金流','百万元','现金流','flow'),'cash':('现金及受限现金','百万元','现金流','stock'),
 'debt':('债务','百万元','现金流','stock'),'gmv':('零售GMV','百万元','零售','flow'),'members':('会员数','百万人','零售','stock'),
 'hotel_cost':('酒店成本','百万元','财务','flow'),'retail_cost':('零售成本','百万元','财务','flow'),'sales_expense':('销售费用','百万元','财务','flow'),'admin_expense':('管理费用','百万元','财务','flow')}
METRICS.update({'monetary_funds':('货币资金','百万元','财务','stock'),'cash_equivalents':('现金及现金等价物','百万元','财务','stock'),'assets':('资产总额','百万元','财务','stock'),'liabilities':('负债总额','百万元','财务','stock')})

def clean(v):
 if isinstance(v,(datetime.date,datetime.datetime)):return v.isoformat()
 return v
def period(v):
 s=str(v).strip()
 if re.fullmatch(r'20\d\d',s):return s+'FY'
 if re.fullmatch(r'20\d\d[QH][1-4]',s):return s
 if re.fullmatch(r'20\d\dA',s):return s[:4]+'FY'
 m=re.fullmatch(r'([1-4])([QH])(\d{2})',s,re.I)
 if m:return '20'+m[3]+m[2].upper()+m[1]
 if re.fullmatch(r'20\d\d[eE]',s):return s[:4]+'FY'
 return None
def numeric(v):
 if isinstance(v,bool):return None
 try:
  n=float(str(v).replace(',','').replace('%',''))
  return n if math.isfinite(n) else None
 except (ValueError,TypeError):return None
def dump(p,v):p.parent.mkdir(parents=True,exist_ok=True);p.write_text(json.dumps(v,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
def write_company(dest,cid,obs):
 parts=[];batch=[];total=0
 for o in obs:
  size=len(json.dumps(o,ensure_ascii=False,separators=(',',':')).encode('utf-8'))+2
  if batch and total+size>35000:
   name=f'hotels/{cid}/{len(parts)+1:03d}.json';parts.append('data/companies/'+name);p=dest/name;p.parent.mkdir(parents=True,exist_ok=True);p.write_text(json.dumps({'observations':batch},ensure_ascii=False,separators=(',',':'))+'\n',encoding='utf-8');batch=[];total=0
  batch.append(o);total+=size
 if batch:
  name=f'hotels/{cid}/{len(parts)+1:03d}.json';parts.append('data/companies/'+name);p=dest/name;p.parent.mkdir(parents=True,exist_ok=True);p.write_text(json.dumps({'observations':batch},ensure_ascii=False,separators=(',',':'))+'\n',encoding='utf-8')
 dump(dest/'hotels'/f'{cid}.json',{'schemaVersion':1,'companyId':cid,'observationCount':len(obs),'parts':parts})
def periodkey(p):return (int(p[:4]),int(p[-1])*3 if 'Q' in p else int(p[-1])*6 if 'H' in p else 12, {'季度':0,'半年':1,'年度':2}.get('季度' if 'Q' in p else '半年' if 'H' in p else '年度'))
def main(source_root,repo):
 root=Path(source_root);dest=Path(repo)/'data/companies';sources={};issues=[];allobs=[]
 def src(rel):
  p=root/rel;sid='src-'+hashlib.sha256(rel.encode()).hexdigest()[:12]
  if sid not in sources:sources[sid]={'id':sid,'file':p.name,'relativePath':rel,'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'kind':'用户整理底稿','verification':'逐单元格转录，未逐页重核原始披露','importedAt':datetime.datetime.now(datetime.timezone(datetime.timedelta(hours=8))).isoformat(timespec='seconds')}
  return sid
 def add(company,p,m,v,sid,locator,scope='全部',mode='全部',region='全部',basis='',status='transcribed',unit=None,currency=None,report=None,pages=None,formula=None,note='',raw=None):
  n=numeric(v)
  if n is None:return
  freq='季度' if 'Q' in p else '半年' if 'H' in p else '年度'
  temporal=METRICS[m][3]
  if temporal=='stock':temporal='期末'
  elif temporal=='average':temporal='期间平均'
  else:temporal='单季' if freq=='季度' else '半年' if freq=='半年' else '全年累计'
  if not unit:unit=METRICS[m][1]
  if not currency:currency='CNY' if unit in ['元/间','百万元'] else None
  ident='|'.join(map(str,[company,p,m,scope,mode,region,basis,sid,locator]))
  o={'id':'obs-'+hashlib.sha256(ident.encode()).hexdigest()[:16],'companyId':company,'period':p,'frequency':freq,'periodBasis':temporal,'metric':m,'value':n,'unit':unit,'currency':currency,'scope':scope,'mode':mode,'region':region,'basis':basis,'status':status,'sourceId':sid,'locator':locator,'report':report,'pages':pages,'formula':formula,'note':note,'rawValue':clean(v if raw is None else raw)}
  if (p[:4]>'2026' or (p.startswith('2026') and p not in ['2026Q1','2026Q2','2026H1'])) and status!='forecast':
   o['status']='pending';o['note']+='；超过已收录实际披露截止期，待核验'
  if m=='occ' and not 0<=n<=100:o['status']='pending';issues.append({'type':'range','observationId':o['id'],'message':'入住率范围异常'})
  allobs.append(o)
  return o
 def xlsx(rel):return openpyxl.load_workbook(root/rel,data_only=True,read_only=False)
 def table(s,header):
  headers=[c.value for c in s[header]]
  for row in s.iter_rows(min_row=header+1):
   if row[0].value is None:continue
   yield row[0].row,{str(k):c.value for k,c in zip(headers,row) if k is not None}
 def csvrows(rel):
  with (root/rel).open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))
 # H World: group stock and China operating data are distinct series.
 rel='AI研报/华住/AI产物/华住经营数据底稿_定期报告页码与真实截图嵌入.xlsx';sid=src(rel);w=xlsx(rel)
 idx=[r for _,r in table(w['截图索引'],1)]
 for sheet,header,mapping in [('季度经营数据',4,{'B':'hotels','C':'rooms','D':'hotels','E':'adr','F':'occ','G':'revpar','H':'revenue','I':'leased_revenue','J':'managed_revenue','K':'other_revenue','M':'operating_profit','N':'parent_profit','O':'adjusted_ebitda'}),('年度数据',3,{'B':'hotels','C':'rooms','D':'adr','E':'occ','F':'revpar','G':'revenue','I':'parent_profit','J':'operating_profit','K':'leased_revenue','L':'managed_revenue','N':'adjusted_ebitda'})]:
  s=w[sheet];note=str(s['A2'].value)
  for cells in s.iter_rows(min_row=header+1):
   p=period(cells[0].value)
   if not p:continue
   evidence=[r for r in idx if str(r['报告期']) in [p,p[:4],p[:4]+'年','FY'+p[:4]]]
   for col,m in mapping.items():
    c=s[f'{col}{cells[0].row}'];oper=m in ['adr','occ','revpar'];china=sheet=='季度经营数据' and ((p[:4]>'2022') or p=='2022Q4') and oper
    add('hworld',p,m,c.value,sid,f'{sheet}!{c.coordinate}',region='中国' if china or (sheet=='季度经营数据' and col=='D') else '集团整体',basis='HWC / Legacy-Huazhu' if china else '集团Blended / 源表年度口径' if oper else '集团整体' if col!='D' else '中国酒店',report=' / '.join(dict.fromkeys(str(r['原始报告']) for r in evidence)),pages=' / '.join(dict.fromkeys(str(r['PDF页码']) for r in evidence)),note=note)
 w.close()
 # Atour: continuous detailed workbook has priority; older workbook extends IPO history.
 rel='AI研报/亚朵/AI产物/亚朵连续数据PDF底稿.xlsx';sid=src(rel);w=xlsx(rel)
 evidence={r['证据ID']:r for _,r in table(w['来源索引'],6)}
 sets=[('年度经营数据',{'B':'hotels','C':'hotels','D':'hotels','E':'rooms','F':'pipeline','G':'occ','H':'adr','I':'revpar'},'J'),('季度经营数据',{'B':'hotels','C':'hotels','D':'hotels','E':'rooms','F':'pipeline','G':'occ','H':'adr','I':'revpar'},'L'),('年度财务数据',{'B':'revenue','C':'managed_revenue','D':'leased_revenue','E':'retail_revenue','F':'other_revenue','G':'retail_other_revenue','H':'operating_profit','K':'net_profit','L':'adjusted_profit','N':'adjusted_ebitda'},'O'),('季度财务数据',{'B':'revenue','C':'managed_revenue','D':'leased_revenue','E':'retail_revenue','F':'other_revenue'},'H'),('现金流与流动性',{'B':'operating_cf','C':'investing_cf','D':'financing_cf','E':'cash','F':'debt'},'G')]
 for sheet,mapping,ecol in sets:
  s=w[sheet];note=str(s['A3'].value)
  for row in s.iter_rows(min_row=7,max_row=12 if sheet.startswith('年度') else 21):
   p=period(row[0].value)
   if not p:continue
   er=evidence.get(s[f'{ecol}{row[0].row}'].value,{})
   for col,m in mapping.items():
    c=s[f'{col}{row[0].row}'];v=c.value
    if m=='occ' and numeric(v) is not None:v=numeric(v)*100
    mode='加盟' if '经营数据' in sheet and col=='C' else '直营' if '经营数据' in sheet and col=='D' else '全部'
    if sheet=='季度财务数据' and col=='E' and p<'2023Q4':m='retail_other_revenue'
    add('atour',p,m,v,sid,f'{sheet}!{c.coordinate}',mode=mode,basis='全酒店 / 含税 / 剔除临时关闭客房' if m in ['adr','occ','revpar'] else '集团整体',report=er.get('原始报告名称'),pages=er.get('PDF物理页码'),note=note,raw=c.value)
   if sheet=='季度经营数据':
    text=str(s[f'K{row[0].row}'].value or '');mt=re.fullmatch(r'\s*(\d+)\s*/\s*(\d+)\s*',text)
    if mt:
     for m,v in zip(['opened','closed'],mt.groups()):add('atour',p,m,v,sid,f'{sheet}!K{row[0].row}',report=er.get('原始报告名称'),pages=er.get('PDF物理页码'),note=note,raw=text)
    elif text:issues.append({'type':'period','sourceId':sid,'locator':f'{sheet}!K{row[0].row}','message':'开关店期间非当季，保留原文：'+text})
    add('atour',p,'gmv',s[f'J{row[0].row}'].value,sid,f'{sheet}!J{row[0].row}',report=er.get('原始报告名称'),pages=er.get('PDF物理页码'),note='零售GMV，区别于收入')
 extras=[('季度财务数据',30,44,{'B':'operating_profit','C':'net_profit','D':'adjusted_profit','F':'adjusted_ebitda'},'H'),('年度财务数据',34,38,{'B':'operating_cf','C':'investing_cf','D':'financing_cf','F':'cash'},'G'),('年度财务数据',21,25,{'B':'hotel_cost','C':'retail_cost','E':'sales_expense','F':'admin_expense'},'I'),('年度经营数据',35,40,{'B':'members','C':'gmv'},'F')]
 for sheet,start,end,mapping,ecol in extras:
  s=w[sheet]
  for row in s.iter_rows(min_row=start,max_row=end):
   p=period(row[0].value)
   if not p:continue
   er=evidence.get(s[f'{ecol}{row[0].row}'].value,{})
   for col,m in mapping.items():
    c=s[f'{col}{row[0].row}'];add('atour',p,m,c.value,sid,f'{sheet}!{c.coordinate}',basis='集团整体',report=er.get('原始报告名称'),pages=er.get('PDF物理页码'),note='原底稿对应子表；GMV不替代收入')
 s=w['年度经营数据']
 for row in s.iter_rows(min_row=21,max_row=26):
  p=period(row[0].value);er=evidence.get(row[7].value,{})
  if not p:continue
  for j,scope in [(1,'开发阶段'),(3,'爬坡阶段'),(5,'成熟阶段')]:
   c=row[j];add('atour',p,'hotels',c.value,sid,f'年度经营数据!{c.coordinate}',scope=scope,basis='生命周期结构',report=er.get('原始报告名称'),pages=er.get('PDF物理页码'))
 s=w['品牌与同店']
 for row in s.iter_rows(min_row=7,max_row=30):
  p=period(row[0].value)
  if not p:continue
  er=evidence.get(row[5].value,{})
  for c,mode in [(row[2],'加盟'),(row[3],'直营')]:add('atour',p,'hotels',c.value,sid,f'品牌与同店!{c.coordinate}',scope=str(row[1].value),mode=mode,basis='品牌酒店',report=er.get('原始报告名称'),pages=er.get('PDF物理页码'),note=str(row[4].value or ''))
 for row in s.iter_rows(min_row=40,max_row=42):
  group=str(row[0].value or '');years=re.findall(r'20\d\d',group)
  if len(years)!=2:continue
  er=evidence.get(row[8].value,{})
  for m,cols in [('occ',[2,3]),('adr',[4,5]),('revpar',[6,7])]:
   for y,j in zip(years,cols):
    c=row[j];add('atour',y+'FY',m,numeric(c.value)*100 if m=='occ' else c.value,sid,f'品牌与同店!{c.coordinate}',scope='同店组 '+group,basis='固定同店组 '+str(row[1].value)+' 家',report=er.get('原始报告名称'),pages=er.get('PDF物理页码'),raw=c.value)
 w.close()
 rel='AI研报/亚朵/AI产物/亚朵酒店经营指标.xlsx';sid=src(rel);w=xlsx(rel)
 existing={(o['period'],o['metric'],o['mode']) for o in allobs if o['companyId']=='atour' and o['scope']=='全部'}
 for sheet,header,mapping in [('季度经营数据',4,{'B':'hotels','C':'rooms','D':'adr','E':'occ','F':'revpar','G':'managed_revenue','H':'leased_revenue','I':'retail_revenue','J':'other_revenue','K':'revenue','L':'net_profit','M':'adjusted_profit','N':'adjusted_ebitda'}),('年度数据',3,{'B':'hotels','C':'rooms','D':'adr','E':'occ','F':'revpar','G':'revenue','I':'net_profit','J':'operating_profit','K':'retail_revenue'})]:
  s=w[sheet]
  for row in s.iter_rows(min_row=header+1,max_row=34 if sheet=='季度经营数据' else 20):
   p=period(row[0].value)
   if not p:continue
   for col,m in mapping.items():
    c=s[f'{col}{row[0].row}']
    if m=='retail_revenue' and p<'2023FY' and sheet=='年度数据':m='retail_other_revenue'
    if (p,m,'全部') in existing:
     prev=next((o for o in allobs if o['companyId']=='atour' and (o['period'],o['metric'],o['mode'],o['scope'])==(p,m,'全部','全部')),None)
     nv=numeric(c.value)
     if prev and nv is not None and abs(prev['value']-nv)>max(0.05,abs(prev['value'])*0.001):issues.append({'type':'conflict','companyId':'atour','period':p,'metric':m,'selectedId':prev['id'],'alternative':{'value':nv,'sourceId':sid,'locator':f'{sheet}!{c.coordinate}'},'message':'两份底稿值不同；主序列采用连续数据PDF底稿，差异保留待核验'})
     continue
    add('atour',p,m,c.value,sid,f'{sheet}!{c.coordinate}',basis='全酒店 / 含税 / 剔除临时关闭客房' if m in ['adr','occ','revpar'] else '集团整体',report=str(s[f'O{row[0].row}'].value or '') if sheet=='季度经营数据' else None,note=str(s['A2'].value))
 w.close()
 # Jin Jiang: source CSV includes physical report pages and computed-quarter labels.
 rel='AI研报/锦江酒店/AI产物/锦江酒店_季度经营指标_2019Q1-2025Q4.csv';sid=src(rel)
 for i,r in enumerate(csvrows(rel),2):
  p=r['年度']+r['季度'];curr='EUR' if '欧元' in r['币种'] else 'CNY'
  for field,m in [('ADR','adr'),('OCC','occ'),('RevPAR','revpar')]:add('jinjiang',p,m,r[field],sid,f'CSV第{i}行/{field}',scope=r['标准分类'],region=r['市场'],basis='全服务型酒店' if '全服务型' in r['标准分类'] else '有限服务型酒店',unit='欧元/间' if m!='occ' and curr=='EUR' else None,currency=curr if m!='occ' else None,report=r['来源报告'],pages=r['PDF物理页码'],note=r['备注'])
 rel='AI研报/锦江酒店/AI产物/锦江酒店_开店情况_季度主数据_2021Q1-2026Q2_UTF8BOM.csv';sid=src(rel)
 for i,r in enumerate(csvrows(rel),2):
  for field,m in [('新开业酒店数','opened'),('开业退出酒店数','closed'),('开业转筹建酒店数','to_pipeline'),('净增开业_计算','net_opened'),('期末开业酒店数','hotels'),('期末开业客房数','rooms')]:add('jinjiang',r['年度']+r['季度'],m,r[field],sid,f'CSV第{i}行/{field}',region='全球',basis=r['口径阶段'],status='derived' if '推算' in r['数据属性'] or m=='net_opened' else 'transcribed',report=r['来源报告'],pages=r['PDF物理页码'],formula=r['计算公式/期间说明'],note=r['备注'])
 rel='AI研报/锦江酒店/AI产物/锦江酒店_开店情况_期末分档次规模_2021Q1-2026Q2_UTF8BOM.csv';sid=src(rel)
 for i,r in enumerate(csvrows(rel),2):
  for field,m in [('期末开业酒店数','hotels'),('期末开业客房数','rooms')]:add('jinjiang',r['年度']+r['季度'],m,r[field],sid,f'CSV第{i}行/{field}',scope=r['分类'] if r['分类维度']=='档次' else '全部',mode=r['分类'] if r['分类维度']!='档次' else '全部',region='全球',basis='有限服务型' if r['年度']<'2023' or r['年度']=='2023' and r['季度']<'Q4' else '酒店业务含全服务型',report=r['来源报告'],pages=r['PDF物理页码'],note=r['备注'])
 rel='AI研报/锦江酒店/AI产物/锦江酒店_境内外分部利润_半年口径_UTF8BOM.csv';sid=src(rel)
 for i,r in enumerate(csvrows(rel),2):
  for field,m,cfield in [('收入_万元','revenue','收入币种'),('归母净利润_万元','parent_profit','利润币种')]:
   n=numeric(r[field]);curr='EUR' if '欧元' in r[cfield] else 'CNY'
   add('jinjiang',r['报告期'],m,n/100 if n is not None else None,sid,f'CSV第{i}行/{field}',region=r['区域标准'],basis=r['业务类型'],currency=curr,status='derived' if r['半年']=='H2' else 'transcribed',report=r['实际来源报告'],pages=r['PDF页码'],formula=r['计算公式'],note=r['数据状态'],raw=r[field])
 # Models: whitelist historical operating rows; forecasts retain a separate status.
 for company,rel in [('btg','酒店/首旅酒店/20260830 首旅酒店模型.xlsx'),('jinjiang','酒店/锦江酒店/20260923 锦江酒店模型.xlsx')]:
  sid=src(rel);w=xlsx(rel);fw=openpyxl.load_workbook(root/rel,data_only=False,read_only=False)
  s=w['利润表'];fs=fw['利润表']
  for c in s[4]:
   p=period(c.value)
   if not p:continue
   forecast='E' in str(c.value).upper()
   if forecast and p[:4]>'2028':continue
   for r,m in [(5,'revenue'),(61,'operating_profit'),(78,'parent_profit')]:
    cc=s.cell(r,c.column);form=fs.cell(r,c.column).value;status='forecast' if forecast else 'transcribed'
    o=add(company,p,m,cc.value,sid,f'利润表!{cc.coordinate}',region='集团整体',basis='模型累计财务',status=status,formula=form if isinstance(form,str) and form.startswith('=') else None,note='模型缓存值，未重算；历史值未逐页复核')
    if o and 'Q' in p:o['periodBasis']='年初至今累计'
  if company=='btg':
   s=w['经营假设'];fs=fw['经营假设'];rows={}
   for base,m in [(50,'hotels'),(120,'rooms'),(160,'revpar'),(181,'adr'),(199,'occ')]:
    offsets=[(0,'经济型','全部'),(2 if base==160 else 1,'经济型','直营'),(3 if base in [50,160] else 2,'经济型','加盟')]
    if base==50:offsets += [(5,'中高端','全部'),(6,'中高端','直营'),(8,'中高端','加盟'),(10,'轻管理','全部'),(17,'全部','全部'),(18,'全部','直营'),(19,'全部','加盟')]
    elif base==120:offsets += [(3,'中高端','全部'),(4,'中高端','直营'),(5,'中高端','加盟'),(6,'轻管理','全部'),(12,'全部','全部'),(13,'全部','直营'),(14,'全部','加盟')]
    elif base==160:offsets += [(4,'中高端','全部'),(6,'中高端','直营'),(7,'中高端','加盟'),(8,'轻管理','全部'),(15,'全部','全部')]
    elif base==181:offsets += [(3,'中高端','全部'),(4,'中高端','直营'),(5,'中高端','加盟'),(6,'轻管理','全部'),(12,'全部','全部')]
    else:offsets += [(3,'中高端','全部'),(4,'中高端','直营'),(5,'中高端','加盟'),(6,'轻管理','全部'),(12,'全部','全部')]
    for off,scope,mode in offsets:rows[base+off]=(m,scope,mode)
   rows.update({70:('opened','全部','全部'),71:('closed','全部','全部'),72:('net_opened','全部','全部'),73:('pipeline','全部','全部')})
   for col in range(2,61):
    rawp=s.cell(48,col).value;p=period(rawp)
    if not p:continue
    forecast='E' in str(rawp).upper()
    for r,(m,scope,mode) in rows.items():
     cc=s.cell(r,col);v=cc.value;form=fs.cell(r,col).value
     if m=='occ' and numeric(v) is not None and abs(numeric(v))<=1:v=numeric(v)*100
     add(company,p,m,v,sid,f'经营假设!{cc.coordinate}',scope=scope,mode=mode,region='中国',basis='模型经营数据 / 含轻管理' if scope=='全部' else '模型经营数据',status='forecast' if forecast else 'transcribed',formula=form if isinstance(form,str) and form.startswith('=') else None,note='模型缓存值；口径见源行标签，待原始报告复核',raw=cc.value)
  w.close();fw.close()
 # Identity and data health are independent of quote dates.
 pool=list(csv.DictReader((Path(repo)/'data/商社-标的池与估值跟踪.csv').open(encoding='utf-8-sig')))
 ids={'600754.SH':'jinjiang','600258.SH':'btg','1179.HK':'hworld','ATAT.O':'atour'}
 companies=[]
 for r in pool:
  if r['子行业']!='酒店' or r['证券代码'] not in ids:continue
  cid=ids[r['证券代码']];obs=[o for o in allobs if o['companyId']==cid]
  actual=[o for o in obs if o['status'] not in ['pending','forecast']]
  companies.append({'id':cid,'name':r['公司名称'],'code':r['证券代码'],'aliases':['HTHT.US','HTHT.O'] if cid=='hworld' else ['ATAT.US'] if cid=='atour' else [],'sector':'酒店','market':r['市场'],'coverage':r['覆盖级别'],'file':f'data/companies/hotels/{cid}.json' if obs else None,'count':len(obs),'firstPeriod':min((o['period'] for o in actual),key=periodkey,default=None),'lastPeriod':max((o['period'] for o in actual),key=periodkey,default=None),'status':'已入库' if obs else '暂无专项经营底稿','valuationDate':r['数据日期'][:10]})
  write_company(dest,cid,obs)
 # Recognized discrepancies are preserved, never fixed by inventing values.
 groups=collections.defaultdict(dict)
 for o in allobs:
  if o['metric'] in ['adr','occ','revpar'] and o['status'] not in ['forecast','pending']:groups[(o['companyId'],o['period'],o['scope'],o['mode'],o['region'],o['basis'])][o['metric']]=o
 for key,g in groups.items():
  if len(g)==3:
   expected=g['adr']['value']*g['occ']['value']/100;actual=g['revpar']['value']
   if abs(actual-expected)>max(1.5,abs(actual)*0.02):issues.append({'type':'reconciliation','companyId':key[0],'period':key[1],'scope':key[2],'observationIds':[o['id'] for o in g.values()],'message':'RevPAR与ADR×OCC差异超过2%或1.5货币单位，保留源值待核验','difference':round(actual-expected,4)})
 dump(dest/'catalog.json',{'schemaVersion':1,'title':'重点公司数据库','updatedAt':'2026-10-10','companies':companies,'metrics':[{'id':k,'name':v[0],'unit':v[1],'category':v[2],'aggregation':v[3]} for k,v in METRICS.items()],'sourcesFile':'data/companies/sources.json','issuesFile':'data/companies/issues.json'})
 dump(dest/'sources.json',{'schemaVersion':1,'sources':list(sources.values())})
 dump(dest/'issues.json',{'schemaVersion':1,'issues':issues,'pendingCount':sum(o['status']=='pending' for o in allobs)})
 fields=['id','companyId','period','frequency','periodBasis','metric','value','unit','currency','scope','mode','region','basis','status','sourceId','locator','report','pages','formula','note','rawValue']
 with (dest/'hotel-observations.csv').open('w',encoding='utf-8-sig',newline='') as f:
  writer=csv.DictWriter(f,fieldnames=fields);writer.writeheader();writer.writerows({k:("'"+v if isinstance(v,str) and v.startswith(('=','+','-','@')) else v) for k,v in o.items()} for o in allobs)
 print(json.dumps({'companies':companies,'sources':len(sources),'observations':len(allobs),'issues':len(issues),'statuses':dict(collections.Counter(o['status'] for o in allobs))},ensure_ascii=False))
if __name__=='__main__':
 ap=argparse.ArgumentParser();ap.add_argument('--source-root',required=True);ap.add_argument('--root',default='.');args=ap.parse_args();main(args.source_root,args.root)
