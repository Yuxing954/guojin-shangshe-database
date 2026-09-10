import csv, re, sys
from pathlib import Path
from openpyxl import load_workbook

src=Path(sys.argv[1]); out=Path(sys.argv[2]); out.mkdir(parents=True,exist_ok=True)
wb=load_workbook(src,data_only=True,read_only=True)

def period(v):
    s=str(v or '').strip(); m=re.search(r'(20\d{2})年第(\d+)周\((20\d{2})-(\d{2})-(\d{2})\s*-\s*(20\d{2})-(\d{2})-(\d{2})\)',s)
    if not m:return None
    return [f'{m[1]}W{int(m[2]):02d}',int(m[1]),int(m[2]),f'{m[3]}-{m[4]}-{m[5]}',f'{m[6]}-{m[7]}-{m[8]}']

def write(name,head,rows):
    p=out/name
    with p.open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.writer(f);w.writerow(head);w.writerows(rows)
    print(f'{name}: {len(rows)}')

# 全国与上海周度经营
industry=[]
for sheet,region_col in [('原始数据(酒店之家)',1),('上海市',1)]:
    ws=wb[sheet]
    for row in list(ws.iter_rows(values_only=True))[1:]:
        p=period(row[0]);
        if not p:continue
        occ=row[7]; adr=row[9] if row[9] is not None else row[8]; rev=row[10]
        if rev in (None,'') and occ not in (None,'') and adr not in (None,''):rev=float(occ)*float(adr)
        industry.append(p+[row[region_col],row[2],row[3],row[4],row[5],row[6],occ,adr,rev,'酒店之家'])
write('hotel_industry_weekly.csv',['period_id','year','week','start_date','end_date','region','segment','hotel_count','room_count','hotel_count_15plus','room_count_15plus','occupancy_rate','adr','revpar','source'],industry)

# 集团经营
groups=[]
for sheet in ['华住','首旅如家','锦江酒店（中国区）','亚朵']:
    ws=wb[sheet]
    for row in list(ws.iter_rows(values_only=True))[1:]:
        p=period(row[0]);
        if not p:continue
        groups.append(p+[row[1],row[2],row[3],row[4],row[5],row[6],row[7],row[8],row[9],row[10],'酒店之家'])
write('hotel_group_weekly.csv',['period_id','year','week','start_date','end_date','group','region','hotel_count','room_count','price_index','heat_index','stay_price_index','stay_heat_index','adr','stay_adr','source'],groups)

# 房量结构
supply=[];ws=wb['分房量']
for row in list(ws.iter_rows(values_only=True))[1:]:
    p=period(row[0]);
    if not p:continue
    supply.append(p+[row[1],row[2],row[3],row[4],row[5],row[6],'酒店之家'])
write('hotel_supply_weekly.csv',['period_id','year','week','start_date','end_date','region','room_band','hotel_count','room_count','chain_hotel_count','chain_room_count','source'],supply)

# 六大集团年度市占率
ws=wb['市占率']; names=[str(ws.cell(3,c).value).split(':')[-1] for c in range(2,8)]; share={}
for row in list(ws.iter_rows(min_row=4,values_only=True)):
    for base in (0,9):
        d=row[base] if base<len(row) else None
        if not d or not re.match(r'20\d{2}',str(d)):continue
        y=int(str(d)[:4])
        for j,name in enumerate(names):
            idx=base+1+j
            if idx<len(row) and row[idx] is not None:share[(y,name)]=row[idx]
write('hotel_market_share_annual.csv',['year','group','market_share_pct','source'],[[y,n,v,'酒店之家/iFinD'] for (y,n),v in sorted(share.items())])

