"""Refresh reviewed Choice series and official history; credentials never enter snapshots."""
import argparse, calendar, hashlib, io, json, math, os, re, sys, tomllib, time, tempfile
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date, datetime, timezone
from pathlib import Path
from urllib.parse import urlparse, parse_qs
from urllib.request import Request, build_opener
from miaoxiang_mcp import Client, URL
from gold_jewelry import NoRedirect, atomic_json

ROOT=Path(__file__).resolve().parents[1]
DATA=ROOT/'data/macro'
class RateLimitError(ValueError):pass

def compact_json(path,data):
    path.parent.mkdir(parents=True,exist_ok=True)
    fd,name=tempfile.mkstemp(prefix='.macro-',dir=path.parent)
    try:
        with os.fdopen(fd,'w',encoding='utf-8') as f:f.write(json.dumps(data,ensure_ascii=False,separators=(',',':'),allow_nan=False)+'\n')
        os.replace(name,path)
    finally:
        if os.path.exists(name):os.unlink(name)

def numeric(value):
    if value in (None,'','-','--','—'):return None
    if isinstance(value,bool):raise ValueError('Boolean numeric value')
    text=str(value).replace(',','').strip()
    match=re.fullmatch(r'(-?\d+(?:\.\d+)?)(万亿|亿|万)?',text)
    if not match:raise ValueError('Unknown numeric representation')
    result=float(match[1])*{None:1,'万':1e4,'亿':1e8,'万亿':1e12}[match[2]]
    if not math.isfinite(result):raise ValueError('Nonfinite value')
    return result, bool(match[2])

def body_of(receipt):
    if receipt.get('isError'):raise ValueError('Provider business error')
    if 'data' in receipt:return receipt
    tables=[]
    for item in receipt.get('content',[]):
        if item.get('type')=='text':
            body=json.loads(item['text'])
            if isinstance(body,dict) and ('操作过于频繁' in str(body.get('message','')) or 'rate limit' in str(body.get('message','')).lower()):raise RateLimitError('Provider rate limit')
            if isinstance(body,dict):tables.extend(body.get('data',[]))
    return {'data':tables}

def valid_period(period,frequency):
    if frequency=='monthly':return bool(re.fullmatch(r'\d{4}-(0[1-9]|1[0-2])',period))
    if frequency=='quarterly':return bool(re.fullmatch(r'\d{4}-Q[1-4]',period))
    if frequency=='daily':
        try:return bool(re.fullmatch(r'\d{4}-\d{2}-\d{2}',period)) and bool(date.fromisoformat(period))
        except ValueError:return False
    return False

def end_of(period,frequency):
    if frequency=='daily':return period
    year=int(period[:4]);month=int(period[-1])*3 if frequency=='quarterly' else int(period[-2:])
    return date(year,month,calendar.monthrange(year,month)[1]).isoformat()

def coverage(observations,spec):
    periods=sorted(r['period'] for r in observations);missing=[];structural=[]
    if periods and spec['frequency'] in ('monthly','quarterly'):
        step=1 if spec['frequency']=='monthly' else 3
        start=int(periods[0][:4])*12+(int(periods[0][-2:])-1 if step==1 else (int(periods[0][-1])-1)*3)
        stop=int(periods[-1][:4])*12+(int(periods[-1][-2:])-1 if step==1 else (int(periods[-1][-1])-1)*3)
        present=set(periods)
        for index in range(start,stop+1,step):
            year,month=divmod(index,12)
            period=f'{year}-{month+1:02}' if step==1 else f'{year}-Q{month//3+1}'
            if period not in present:
                (structural if month+1 in spec.get('expectedMissingMonths',[]) else missing).append(period)
    requested=spec.get('start')
    requested_period=(requested[:7] if spec['frequency']=='monthly' else requested[:4]+'-Q'+str((int(requested[5:7])-1)//3+1)) if requested and spec['frequency']!='daily' else None
    # Non-publication at the beginning of a year is not truncated history.
    earliest=next((p for p in periods if spec['frequency']!='monthly' or int(p[-2:]) not in spec.get('expectedMissingMonths',[])),None)
    expected=requested_period
    if expected and spec['frequency']=='monthly':
        while int(expected[-2:]) in spec.get('expectedMissingMonths',[]):
            y,m=int(expected[:4]),int(expected[-2:]);expected=f'{y+1 if m==12 else y}-{1 if m==12 else m+1:02}'
    short=bool(expected and earliest and earliest>expected)
    return {'start':periods[0] if periods else None,'end':periods[-1] if periods else None,'count':len(periods),'requestedStart':requested,'historyShort':short,
            'missingPeriods':missing,'structuralPeriods':structural,'complete':not missing and not short,
            'calendar': '交易日序列；不将周末及节假日算作缺失' if spec['frequency']=='daily' else '原始发布频率'}

def normalize(receipt,spec,cutoff):
    observations={};links=set();codes=set();matched=0
    for table in body_of(receipt).get('data',[]):
        columns=table.get('columns',[])
        # A semantic search may also return annual or otherwise incompatible tables.
        periods=columns[2:]
        if not periods or not all(valid_period(p,spec['frequency']) for p in periods):continue
        table_links=[x.get('choiceUrl','') for x in table.get('meta',{}).get('jumpUrlList',[])]
        table_codes=set()
        for link in table_links:
            u=urlparse(link)
            if u.scheme!='https' or u.hostname!='choicew2z.eastmoney.com':raise ValueError('Untrusted source URL')
            table_codes.update(parse_qs(u.fragment.split('?',1)[-1]).get('macroids',[''])[0].split(','))
        for row in table.get('items',[]):
            if len(row)<2 or row[0]!=spec['originalName'] or row[1]!=spec['publisher']:continue
            if len(row)!=len(columns):raise ValueError('Source dimensions differ')
            if spec.get('code') and spec['code'] not in table_codes:raise ValueError('Indicator code mismatch')
            if spec.get('codeCandidates') and not table_codes.intersection(spec['codeCandidates']):raise ValueError('Indicator code mismatch')
            matched+=1;links.update(table_links);codes.update(table_codes)
            for period,raw in zip(periods,row[2:]):
                if end_of(period,spec['frequency'])>cutoff or end_of(period,spec['frequency'])<spec['start']:continue
                parsed=numeric(raw)
                if parsed is None:continue
                value,rounded=parsed;value=round(value*spec.get('scale',1)+spec.get('offset',0),10)
                if period in observations and observations[period]['value']!=value:raise ValueError('Conflicting observations')
                observations[period]={'period':period,'value':value,'releaseDate':None,'sourceUrl':next(iter(table_links),None),
                    'footnotes':['源接口金额已舍入，保留原精度'] if rounded else []}
    if not matched or not observations:raise ValueError('No exact reviewed indicator')
    result={'status':'ready','observations':sorted(observations.values(),key=lambda r:r['period']),
        'sourceOrganization':spec['publisher'],'provider':'东方财富 Choice / 妙想 MCP','providerCodes':sorted(codes),
        'sourceUrl':sorted(links)[0],'receiptSha256':hashlib.sha256(json.dumps(receipt,ensure_ascii=False,sort_keys=True).encode()).hexdigest(),
        'metric':spec.get('metric'), 'spec':{**{k:spec[k] for k in ('name','unit','kind','frequency','adjustment') if k in spec},'definition':spec.get('adjustment',spec['originalName'])},
        'registryId':spec['id'],'coverage':coverage(list(observations.values()),spec)}
    if spec.get('consumerId'):result.update(consumerId=spec['consumerId'],basis=spec['basis'],consumerMeasure=spec.get('consumerMeasure','value'))
    return result

def merge_history(previous,result,spec,now):
    # Same reviewed metric only: never join an old level with a newly selected growth rate.
    fingerprint=hashlib.sha256(json.dumps({k:spec.get(k) for k in ('id','originalName','publisher','unit','kind','frequency','code','scale','offset')},sort_keys=True).encode()).hexdigest()
    if previous.get('identitySha256') and previous['identitySha256']!=fingerprint:raise ValueError('Reviewed series definition changed')
    old=previous.get('observations',[]) if previous.get('registryId')==spec['id'] else []
    rows={r['period']:r for r in old};revisions=0
    for row in result['observations']:
        existing=rows.get(row['period'])
        if existing and existing['value']!=row['value']:revisions+=1
        if existing and existing['value']==row['value'] and existing.get('releaseDate'):row={**row,'releaseDate':existing['releaseDate']}
        rows[row['period']]=row
    result={**result,'identitySha256':fingerprint,'observations':sorted(rows.values(),key=lambda r:r['period']),'fetchedAt':now,'checkedAt':now,'revisedCount':revisions}
    result['coverage']=coverage(result['observations'],spec)
    return result

def fetch_bytes(url):
    with build_opener(NoRedirect()).open(Request(url,headers={'User-Agent':'GuojinMacro/2.0'}),timeout=60) as response:
        return response.read(12_000_001)

def bea_gdp(cutoff):
    from openpyxl import load_workbook
    url='https://apps.bea.gov/national/Release/XLS/Survey/Section1All_xls.xlsx'
    raw=fetch_bytes(url)
    if len(raw)>12_000_000:raise ValueError('Workbook too large')
    book=load_workbook(io.BytesIO(raw),read_only=True,data_only=True);rows=list(book['T10101-Q'].values)
    if not str(rows[0][0]).startswith('Table 1.1.1. Percent Change From Preceding Period in Real Gross Domestic Product'):raise ValueError('BEA table identity')
    header=next(r for r in rows if r[0]=='Line');target=next(r for r in rows if r[2]=='A191RL')
    observations=[]
    for period,value in zip(header[3:],target[3:]):
        if not isinstance(period,str) or not re.fullmatch(r'\d{4}Q[1-4]',period):continue
        period=period[:4]+'-'+period[4:]
        if period<'2016-Q1' or end_of(period,'quarterly')>cutoff or not isinstance(value,(int,float)):continue
        observations.append({'period':period,'value':float(value),'releaseDate':None,'sourceUrl':url})
    spec={'id':'us-gdp-quarter','frequency':'quarterly'}
    return {'status':'ready','registryId':'bea-real-gdp','observations':observations,'provider':'BEA 官方工作簿',
        'sourceOrganization':'美国经济分析局 BEA','sourceUrl':url,'receiptSha256':hashlib.sha256(raw).hexdigest(),
        'spec':{'name':'实际GDP季度增长（折年率）','unit':'%','kind':'rate','frequency':'quarterly','adjustment':'季调环比折年率'},'coverage':coverage(observations,spec)}

def bls_history(catalog,cutoff):
    from macro_data import fetch,observation,period_end
    wanted=[s for s in catalog['series'] if s['adapter']=='bls'];out={s['id']:[] for s in wanted}
    for year in range(2016,int(cutoff[:4])+1,3):
        payload={'seriesid':[s['code'] for s in wanted],'startyear':str(year),'endyear':str(min(year+2,int(cutoff[:4])))}
        body,digest=fetch('https://api.bls.gov/publicAPI/v2/timeseries/data/',payload)
        if body.get('status')!='REQUEST_SUCCEEDED':raise ValueError('BLS rejected history')
        for series in body['Results']['series']:
            spec=next(s for s in wanted if s['code']==series['seriesID'])
            for row in series['data']:
                if not re.fullmatch(r'M(0[1-9]|1[0-2])',row['period']):continue
                period=row['year']+'-'+row['period'][1:]
                if period_end(period,'monthly').isoformat()>cutoff or row['value'] in ('-',None,''):continue
                out[spec['id']].append(observation(period,float(row['value'])))
    return {s['id']:{'status':'ready','registryId':'bls-'+s['code'],'observations':sorted(out[s['id']],key=lambda r:r['period']),
        'sourceOrganization':'美国劳工统计局 BLS','provider':'BLS 官方 API','sourceUrl':'https://data.bls.gov/timeseries/'+s['code'],
        'coverage':coverage(out[s['id']],s)} for s in wanted if out[s['id']]}

def census_housing(cutoff):
    from openpyxl import load_workbook
    url='https://www.census.gov/construction/nrc/xls/starts_cust.xlsx'
    request=Request(url,headers={'User-Agent':'Mozilla/5.0','Referer':'https://www.census.gov/construction/nrc/data/series.html'})
    with build_opener(NoRedirect()).open(request,timeout=60) as response:raw=response.read(12_000_001)
    if len(raw)>12_000_000:raise ValueError('Workbook too large')
    rows=list(load_workbook(io.BytesIO(raw),read_only=True,data_only=True)['Seasonally Adjusted'].values)
    if rows[0][0]!='New Privately-Owned Housing Units Started' or rows[3][0]!='Seasonally adjusted annual rate' or rows[4][1]!='United States' or rows[5][1]!='Total':raise ValueError('Census table identity')
    observations=[]
    for row in rows[6:]:
        if not isinstance(row[0],datetime) or not isinstance(row[1],(int,float)):continue
        period=row[0].strftime('%Y-%m')
        if period<'2016-01' or end_of(period,'monthly')>cutoff:continue
        observations.append({'period':period,'value':float(row[1]),'releaseDate':None,'sourceUrl':url})
    spec={'frequency':'monthly','start':'2016-01-01'}
    return {'status':'ready','registryId':'census-housing-saar','observations':observations,'provider':'Census 官方工作簿','sourceOrganization':'美国人口普查局 Census','sourceUrl':url,'receiptSha256':hashlib.sha256(raw).hexdigest(),
        'spec':{'name':'新屋开工','unit':'千套','kind':'level','frequency':'monthly','adjustment':'全国总量，季调折年率','definition':'新建私人住宅开工总量；美国人口普查局与HUD联合调查，季调折年率。'},'coverage':coverage(observations,spec)}

def main():
    ap=argparse.ArgumentParser(description=__doc__);ap.add_argument('--codex-config',type=Path);ap.add_argument('--receipts-dir',type=Path)
    ap.add_argument('--cutoff',default=date.today().isoformat());ap.add_argument('--only',nargs='*');ap.add_argument('--skip-official',action='store_true');args=ap.parse_args()
    date.fromisoformat(args.cutoff);registry=json.loads((DATA/'choice-registry.json').read_text(encoding='utf-8'))
    if registry.get('publicRedistributionApproved') is not True:raise ValueError('Publication authorization required')
    target=DATA/'automatic-series.json';previous=json.loads(target.read_text(encoding='utf-8')) if target.exists() else {'series':{}}
    values=dict(previous['series']);now=datetime.now(timezone.utc).isoformat(timespec='seconds');errors=[]
    entries=[s for s in registry['series'] if not args.only or s['id'] in args.only]
    key=os.environ.get('EM_API_KEY','')
    if not key and args.codex_config:
        server=tomllib.loads(args.codex_config.read_text(encoding='utf-8'))['mcp_servers']['mx-ds-mcp']
        if server['url']!=URL:raise ValueError('Unexpected provider endpoint')
        key=server['http_headers']['em_api_key']
    receipts=[]
    if args.receipts_dir:
        for path in args.receipts_dir.glob('choice-*.json'):
            try:receipts.append(json.loads(path.read_text(encoding='utf-8')))
            except ValueError:pass
    remote={'client':None,'receipts':{}}
    def task(spec):
        if receipts:
            candidates=[]
            for r in receipts:
                try:candidates.append(normalize(r,spec,args.cutoff))
                except ValueError:pass
            if candidates:return max(candidates,key=lambda x:len(x['observations']))
        if not key:raise ValueError('Choice credentials missing')
        if remote['client'] is None:
            remote['client']=Client(key);remote['client'].initialize()
        client=remote['client']
        start=spec['start']
        if spec.get('lookbackYears'):start=str(int(args.cutoff[:4])-spec['lookbackYears'])+'-01-01'
        frequency={'monthly':'月度','quarterly':'季度','daily':'日度'}[spec['frequency']]
        query=f"{spec['originalName']}，来源{spec['publisher']}，{frequency}，{start}至{args.cutoff}完整历史，保留原始名称、精度和来源链接。"
        if query in remote['receipts']:return normalize(remote['receipts'][query],spec,args.cutoff)
        for attempt in range(4):
            time.sleep(4 if attempt==0 else 20*attempt)
            result=client.rpc('tools/call',{'name':'mx_macro_data','arguments':{'query':query}})
            try:body_of(result);break
            except RateLimitError:
                if attempt==3:raise
        remote['receipts'][query]=result
        raw=json.dumps(result,ensure_ascii=False).encode();digest=hashlib.sha256(raw).hexdigest()
        receipt_dir=ROOT/'.cache/macro-receipts';receipt_dir.mkdir(parents=True,exist_ok=True)
        (receipt_dir/(digest[:16]+'.json')).write_bytes(raw)
        return normalize(result,spec,args.cutoff)
    # The provider rate-limits a shared key; serialize requests and reuse duplicate queries.
    with ThreadPoolExecutor(max_workers=1) as pool:
        futures={pool.submit(task,s):s for s in entries}
        for future in as_completed(futures):
            spec=futures[future]
            try:values[spec['id']]=merge_history(values.get(spec['id'],{}),future.result(),spec,now)
            except Exception as exc:
                safe_reasons={'No exact reviewed indicator','MCP protocol error','Missing or ambiguous MCP response','Provider business error','Reviewed series definition changed','Indicator code mismatch'}
                errors.append({'id':spec['id'],'errorType':type(exc).__name__,'reason':str(exc) if str(exc) in safe_reasons else '接口或数据校验失败'})
                values[spec['id']]={**values.get(spec['id'],{}),'status':'error','checkedAt':now,'error':'接口返回未通过指标与口径校验，保留已有历史。'}
    if not args.skip_official:
        for name,fn in [('bea-real-gdp',lambda:{'us-gdp-quarter':bea_gdp(args.cutoff)}),('census-housing',lambda:{'us-housing-starts':census_housing(args.cutoff)}),('bls-history',lambda:bls_history(json.loads((DATA/'catalog.json').read_text(encoding='utf-8')),args.cutoff))]:
            try:
                for id,result in fn().items():
                    # Authoritative BLS retains index/seasonal definitions, with full history.
                    spec={'id':result['registryId'],'frequency':result.get('spec',{}).get('frequency','monthly'),'start':'2016-01-01'}
                    prior=values.get(id,{})
                    if prior.get('registryId')==result['registryId']:
                        rows={r['period']:r for r in prior.get('observations',[])};rows.update({r['period']:r for r in result['observations']})
                        result['observations']=sorted(rows.values(),key=lambda r:r['period']);result['coverage']=coverage(result['observations'],spec)
                    values[id]={**result,'fetchedAt':now,'checkedAt':now}
            except Exception as exc:
                errors.append({'id':name,'errorType':type(exc).__name__})
                ids=['us-gdp-quarter'] if name=='bea-real-gdp' else ['us-housing-starts'] if name=='census-housing' else [s['id'] for s in json.loads((DATA/'catalog.json').read_text(encoding='utf-8'))['series'] if s['adapter']=='bls']
                for id in ids:values[id]={**values.get(id,{}),'status':'error','checkedAt':now,'error':'官方接口更新失败，保留同口径历史。'}
    output={'version':1,'cutoff':args.cutoff,'checkedAt':now,'publicRedistributionApproved':True,'series':values,'errors':errors}
    for series in values.values():
        for row in series.get('observations',[]):
            if not row.get('sourceUrl') or row['sourceUrl']==series.get('sourceUrl'):row.pop('sourceUrl',None)
            if not row.get('footnotes'):row.pop('footnotes',None)
    compact_json(target,output)
    atomic_json(DATA/'coverage.json',{'version':1,'checkedAt':now,'series':{k:v.get('coverage',{}) for k,v in values.items()},'errors':errors})
    print(json.dumps({'series':len(values),'ready':sum(bool(v.get('observations')) for v in values.values()),'failures':errors},ensure_ascii=False))
    return 1 if errors else 0

if __name__=='__main__':
    try:sys.exit(main())
    except Exception as exc:print(json.dumps({'ok':False,'errorType':type(exc).__name__}));sys.exit(2)
